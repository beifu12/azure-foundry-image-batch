"""Container UI/API regressions. Synthetic data only; no Azure credentials."""
import io
import json
import threading
import time
import unittest
from urllib.request import Request, build_opener, ProxyHandler
from urllib.error import HTTPError
from unittest.mock import Mock, patch

import web_app as web
import image_batch as batch

urlopen=build_opener(ProxyHandler({})).open

TENANT='33333333-3333-4333-8333-333333333333'
SID='11111111-1111-4111-8111-111111111111'
AID=f'/subscriptions/{SID}/resourceGroups/CV-A03/providers/Microsoft.CognitiveServices/accounts/cv-a03-001'


class MemoryAzure:
    """ARM-shaped state machine; records writes without contacting any cloud."""
    def __init__(self):self.resources={};self.writes=[]
    def all(self,path,params=None):
        if path=='/subscriptions':return [{'subscriptionId':SID,'tenantId':TENANT,'state':'Enabled','displayName':'Azure Subscription A03'}]
        if path.endswith('/locations'):return [{'name':name} for name in ('eastus','eastus2','westus3')]
        if path.endswith('/accounts'):return [v for k,v in self.resources.items() if k.endswith('/'+v.get('name','')) and v.get('kind')=='AIServices']
        if path.endswith('/deployments'):return [v for k,v in self.resources.items() if k.startswith(path+'/')]
        if path.endswith('/projects'):return [v for k,v in self.resources.items() if k.startswith(path+'/')]
        if path.endswith('/modelCapacities'):
            locations=['eastus2','westus3'] if params['modelName']=='gpt-image-2' else ['eastus','eastus2','westus3']
            return [{'location':r,'properties':{'skuName':'GlobalStandard','availableCapacity':1}} for r in locations]
        if path.endswith('/models'):return [{'model':{'format':'OpenAI','name':name,'version':'2026-01-01','isDefaultVersion':True,'skus':[{'name':'GlobalStandard','usageName':name,'capacity':{'minimum':1,'maximum':1,'step':1},'rateLimits':[{'key':'request','count':1,'renewalPeriod':60}]+([] if name.startswith('gpt-image-') else [{'key':'token','count':1000,'renewalPeriod':60}])}]}} for name in batch.MODELS]
        if path.endswith('/usages'):return [{'name':{'value':name},'unit':'Count','currentValue':0,'limit':1} for name in batch.MODELS]
        raise AssertionError(path)
    def optional(self,path,version=None):return self.resources.get(path)
    def get(self,path,version=None):return self.resources[path]
    def put(self,path,body,version=None):
        self.writes.append(path);value=copy_dict(body);value.update(id=path,name=path.rsplit('/',1)[-1]);value.setdefault('properties',{})['provisioningState']='Succeeded'
        if '/accounts/' in path and '/deployments/' not in path and '/projects/' not in path:
            value['identity']={'type':'SystemAssigned','principalId':'synthetic-principal'};value['properties'].update(allowProjectManagement=True,endpoint='https://example.invalid')
        self.resources[path]=value;return value
    def call(self,method,path,body=None,params=None,headers=None):
        if method=='POST' and path.endswith('/listKeys'):return {'key1':'synthetic-key'}
        raise AssertionError((method,path))


def copy_dict(value):return json.loads(json.dumps(value))


class WorkflowTests(unittest.TestCase):
    def test_prefetch_reuses_each_workers_http_connection_and_closes_it(self):
        session=web.Session();session.credential=Mock();transports=[]
        def transport():
            value=Mock();value.trust_env=False;transports.append(value);return value
        def response(*args,**kwargs):time.sleep(.005);return {'value':[]}
        with patch('requests.Session',side_effect=transport),patch.object(batch.Azure,'call',response):
            api=web.job_api(session,threading.Event(),snapshot=True)
            web.prefetch_reads(api,[(f'/catalog/{i}',None) for i in range(20)])
        self.assertLessEqual(len(transports),5)
        for value in transports[1:]:value.close.assert_called_once()

    def test_prefetch_global_limit_applies_across_subscription_workers(self):
        from concurrent.futures import ThreadPoolExecutor
        active=0;maximum=0;lock=threading.Lock()
        def response(api,method,path,*args,**kwargs):
            nonlocal active,maximum
            with lock:active+=1;maximum=max(maximum,active)
            try:time.sleep(.025);return {'value':[]}
            finally:
                with lock:active-=1
        session=web.Session();session.credential=Mock();event=threading.Event()
        with patch.object(batch.Azure,'call',response),ThreadPoolExecutor(max_workers=8) as pool:
            jobs=[pool.submit(web.prefetch_reads,web.job_api(session,event,snapshot=True),[(f'/s{i}/r{j}',None) for j in range(8)]) for i in range(8)]
            for job in jobs:job.result()
        self.assertGreater(maximum,4);self.assertLessEqual(maximum,web.MAX_CHECK_READS)

    def test_failed_second_page_cannot_be_read_as_empty_or_complete_inventory(self):
        def response(api,method,path,*args,**kwargs):
            if path=='/accounts':return {'value':[{'name':'known'}],'nextLink':batch.ARM+'/denied'}
            raise batch.AzureError(403,'AuthorizationFailed')
        session=web.Session();session.credential=Mock();api=web.job_api(session,threading.Event(),snapshot=True)
        with patch.object(batch.Azure,'call',response):
            web.prefetch_reads(api,[('/accounts',None)])
            with self.assertRaises(batch.AzureError):api.all('/accounts')

    def test_merged_plan_retains_each_subscriptions_default_version_for_fresh_scan(self):
        second='22222222-2222-4222-8222-222222222222'
        session=web.Session();session.api=True;session.tenant=TENANT
        session.subscriptions=[{'subscriptionId':sid,'tenantId':TENANT,'state':'Enabled','displayName':f'A0{i}'} for i,sid in enumerate((SID,second),3)]
        data={'subscription_ids':[SID,second],'models':[{'name':'gpt-image-2'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        def scope(_session,configs,template,event,notify,cache=None):
            config=configs[0];sid=config['subscription_ids'][0];config['model_version']='v1' if sid==SID else 'v2'
            config['subscription_versions']={sid:config['model_version']}
            row={'subscription_id':sid,'subscription':'A03','region':'eastus2','model':config['model'],'version':config['model_version'],'status':'planned'}
            plan={'config':config,'tenant_id':TENANT,'rows':[row],'errors':[]}
            return [{'plan':plan,'authoritative':copy_dict(plan),'projects':{},'project_errors':[]}],[]
        captured={}
        def execute(api,expected,**kwargs):
            captured[expected['config']['subscription_ids'][0]]=expected['config']['model_version']
            return {'rows':[]}
        with patch.object(web,'JobAzure',return_value=MemoryAzure()),patch.object(web.batch,'subscriptions',return_value=session.subscriptions),patch.object(web,'scan_scope',side_effect=scope):
            web.start_scan(session,data);session.thread.join(3)
        with patch.object(batch,'execute',side_effect=execute),patch.object(web,'verify_results'),patch.object(web,'collect_keys'):
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(3)
        self.assertEqual(captured,{SID:'v1',second:'v2'})

    def test_model_entry_routes_web_and_legacy_cli_without_duplicating_workflow(self):
        import foundry_models
        with patch.object(web,'main',return_value=0) as start_web:
            self.assertEqual(foundry_models.main(['web','--port','8767']),0)
            start_web.assert_called_once_with(['--port','8767'])
        with patch.object(batch,'main',return_value=0) as start_cli:
            self.assertEqual(foundry_models.main(['doctor','--offline','--auth','app']),0)
            start_cli.assert_called_once_with(['doctor','--offline','--auth','app'])

    def test_existing_deployment_still_checks_quota_and_zero_balance_does_not_create(self):
        api=MemoryAzure();api.put(AID,{'kind':'AIServices','location':'eastus2','properties':{}})
        api.put(AID+'/deployments/image',{'sku':{'name':'GlobalStandard','capacity':1},'properties':{'model':{'format':'OpenAI','name':'gpt-image-2','version':'2026-01-01'}}});api.writes=[]
        original=api.all;reads=[]
        def quota(path,params=None):
            reads.append(path)
            if path.endswith('/usages'):return [{'name':{'value':name},'unit':'Count','currentValue':1,'limit':1} for name in batch.MODELS]
            return original(path,params)
        session=web.Session();session.api=True;session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID,'displayName':'A03'}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        configs,template=web.configurations(session,data)
        with patch.object(web,'JobAzure',return_value=api),patch.object(api,'all',side_effect=quota):
            bundles,_=web.scan_scope(session,configs,template,threading.Event(),lambda *a,**k:None)
        row=bundles[0]['plan']['rows'][0]
        self.assertEqual(row['status'],'existing');self.assertTrue(row['quota_checked']);self.assertEqual(row['quota_remaining'],0)
        self.assertTrue(any(p.endswith('/usages') for p in reads));self.assertEqual(api.writes,[])

    def test_prefetch_bounds_parallel_reads_and_reuses_complete_pages_and_failures(self):
        from collections import Counter
        calls=Counter();active=0;maximum=0;lock=threading.Lock()
        def response(api,method,path,body=None,params=None,headers=None):
            nonlocal active,maximum
            self.assertEqual(method,'GET')
            with lock:calls[path]+=1;active+=1;maximum=max(maximum,active)
            try:
                time.sleep(.01)
                if path.endswith('/denied'):raise batch.AzureError(403,'AuthorizationFailed')
                if path.endswith('/accounts'):return {'value':[{'name':'first'}],'nextLink':batch.ARM+'/page-two'}
                return {'value':[{'name':'second'}]}
            finally:
                with lock:active-=1
        session=web.Session();session.credential=Mock();event=threading.Event()
        with patch.object(batch.Azure,'call',response):
            api=web.job_api(session,event,snapshot=True)
            paths=[('/accounts',None),('/accounts',None),('/models',None),('/usages',None),('/denied',None)]
            web.prefetch_reads(api,paths)
            self.assertEqual([r['name'] for r in api.all('/accounts')],['first','second'])
            for _ in range(2):
                with self.assertRaises(batch.AzureError):api.all('/denied')
            self.assertTrue(all(count==1 for count in calls.values()),calls)
            self.assertGreater(maximum,1);self.assertLessEqual(maximum,web.MAX_CHECK_READS)
            api.all('/accounts')[0]['name']='changed'
            self.assertEqual(api.all('/accounts')[0]['name'],'first')
            # A new check does not inherit either stale data or permission failures.
            fresh=web.job_api(session,event,snapshot=True)
            self.assertEqual(len(fresh.all('/accounts')),2)
            self.assertEqual(calls['/accounts'],2)

    def test_prefetch_is_cancelled_and_never_submits_writes(self):
        session=web.Session();session.credential=Mock();event=threading.Event();event.set()
        api=web.job_api(session,event,snapshot=True)
        with patch.object(batch.Azure,'call',side_effect=AssertionError('cancelled check must not request Azure')):
            with self.assertRaises(web.Cancelled):web.prefetch_reads(api,[('/accounts',None)])

    def test_azure_qualified_project_name_is_exported_and_verified_as_child_name(self):
        api=MemoryAzure();api.put(AID,{'kind':'AIServices','location':'eastus2','properties':{}})
        api.put(AID+'/projects/cv-a03-001-resource',{'properties':{}})
        api.resources[AID+'/projects/cv-a03-001-resource']['name']='cv-a03-001/cv-a03-001-resource'
        api.put(AID+'/deployments/image',{'sku':{'name':'GlobalStandard','capacity':1},'properties':{'model':{'format':'OpenAI','name':'gpt-image-2','version':'v'}}});api.writes=[]
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID,'displayName':'Azure Subscription A03'}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'v'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(3)
            self.assertEqual(session.plans[0]['projects'][AID]['name'],'cv-a03-001-resource')
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertTrue(session.results[0]['verified']);self.assertEqual(api.writes,[])

    def test_partial_existing_read_never_recreates_a_disappeared_project(self):
        api=MemoryAzure();api.put(AID,{'kind':'AIServices','location':'eastus2','properties':{}})
        for name in ('one','two'):api.put(AID+'/deployments/'+name,{'sku':{'name':'GlobalStandard','capacity':1},'properties':{'model':{'format':'OpenAI','name':'gpt-image-2','version':'v'}}})
        api.writes=[];session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.plan_id='reviewed'
        row={'subscription':'A03','subscription_id':SID,'region':'eastus2','model':'gpt-image-2','status':'existing','existing':[{'account_id':AID,'deployment':name,'sku':'GlobalStandard','capacity':1,'version':'v'} for name in ('one','two')]}
        plan={'rows':[row],'errors':[],'config':{'subscription_ids':[SID]},'tenant_id':TENANT}
        session.plans=[{'plan':plan,'projects':{AID:{'name':'p','new':False}}}]
        original=api.get
        def missing(path,version=None):
            if path.endswith('/deployments/two') or '/projects/' in path:raise batch.AzureError(404,'ResourceNotFound')
            return original(path,version)
        with patch.object(web,'JobAzure',return_value=api),patch.object(api,'get',side_effect=missing):
            web.start_deploy(session,{'plan_id':'reviewed','confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'partial',session.job);self.assertEqual(api.writes,[])

    def test_existing_only_result_does_not_rescan_every_subscription_model(self):
        api=MemoryAzure()
        api.put(AID,{'kind':'AIServices','location':'eastus2','properties':{}})
        api.put(AID+'/projects/p',{'properties':{}})
        api.put(AID+'/deployments/image',{'sku':{'name':'GlobalStandard','capacity':1},'properties':{'model':{'format':'OpenAI','name':'gpt-image-2','version':'v'}}})
        api.writes=[]
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.plan_id='reviewed'
        row={'subscription':'A03','subscription_id':SID,'region':'eastus2','model':'gpt-image-2','status':'existing','existing':[{'account_id':AID,'deployment':'image','sku':'GlobalStandard','capacity':1,'version':'v'}]}
        plan={'rows':[row],'errors':[],'config':{'subscription_ids':[SID]},'tenant_id':TENANT}
        session.plans=[{'plan':plan,'projects':{AID:{'name':'p','new':False}}}]
        with patch.object(web,'JobAzure',return_value=api),patch.object(batch,'execute',side_effect=AssertionError('只读复用不应重新扫描全部库存')):
            web.start_deploy(session,{'plan_id':'reviewed','confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertTrue(session.results[0]['verified']);self.assertEqual(len(session.keys),1)
        self.assertEqual(api.writes,[])

    def test_preflight_reuses_successful_inventory_reads_but_execution_reads_stay_fresh(self):
        session=web.Session();session.credential=Mock();event=threading.Event()
        with patch.object(batch.Azure,'call',return_value={'value':[{'name':'synthetic'}]}) as call:
            snapshot=web.job_api(session,event,snapshot=True)
            first=snapshot.call('GET',AID+'/deployments');first['value'].clear()
            self.assertEqual(len(snapshot.call('GET',AID+'/deployments')['value']),1)
            self.assertEqual(call.call_count,1)
            with self.assertRaises(ValueError):snapshot.call('PUT',AID,{'properties':{}})
            live=web.job_api(session,event)
            live.call('GET',AID+'/deployments');live.call('GET',AID+'/deployments')
            self.assertEqual(call.call_count,3)

    def test_fifty_five_subscriptions_are_scanned_deployed_verified_and_exportable(self):
        from uuid import uuid5,NAMESPACE_DNS
        subs=[{'subscriptionId':str(uuid5(NAMESPACE_DNS,'synthetic-sub-'+str(i))),'tenantId':TENANT,'state':'Enabled','displayName':f'Azure Subscription A{i:02d}'} for i in range(1,56)]
        class Many(MemoryAzure):
            def __init__(self):super().__init__();self.mutex=threading.Lock();self.active=0;self.maximum=0
            def all(self,path,params=None):
                if path=='/subscriptions':return subs
                if path.endswith('/accounts'):
                    with self.mutex:self.active+=1;self.maximum=max(self.maximum,self.active)
                    try:
                        time.sleep(.002);prefix=path.split('/providers/')[0]+'/'
                        return [v for k,v in list(self.resources.items()) if k.startswith(prefix) and v.get('kind')=='AIServices']
                    finally:
                        with self.mutex:self.active-=1
                return super().all(path,params)
        api=Many();session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=subs
        data={'subscription_ids':[s['subscriptionId'] for s in subs],'models':[{'name':'gpt-image-2','version':'2026-01-01'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(10)
            self.assertEqual(len(session.plans[0]['plan']['rows']),55)
            self.assertGreater(api.maximum,1);self.assertLessEqual(api.maximum,web.MAX_CHECK_SUBSCRIPTIONS)
            api.maximum=0
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(10)
            self.assertLessEqual(api.maximum,web.MAX_SUBSCRIPTIONS_IN_FLIGHT)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(len(session.results),55);self.assertTrue(all(r['verified'] for r in session.results))
        self.assertEqual(len(session.keys),55)
        self.assertEqual(len(api.writes),len(set(api.writes)))

    def test_final_verification_rejects_capacity_drift_and_never_writes(self):
        api=MemoryAzure()
        api.put(AID,{'kind':'AIServices','location':'eastus2','properties':{}})
        api.put(AID+'/projects/p',{'properties':{}})
        api.put(AID+'/deployments/image',{'sku':{'name':'GlobalStandard','capacity':1},'properties':{'model':{'format':'OpenAI','name':'gpt-image-2','version':'v'}}})
        api.writes=[]
        session=web.Session();session.api=Mock();session.credential=Mock()
        session.results=[{'订阅':'A03','订阅 ID':SID,'资源名称-foundry':'account','资源':'p','_account_id':AID,'model':'gpt-image-2','version':'v','sku':'GlobalStandard','capacity':2,'deployment':'image','verified':False}]
        with patch.object(web,'JobAzure',return_value=api):web.verify_results(session,threading.Event())
        self.assertFalse(session.results[0]['verified'])
        self.assertTrue(session.job['partial'])
        self.assertEqual(api.writes,[])

    def test_uncertain_write_is_not_replayed_after_rescan_in_same_session(self):
        session=web.Session();session.credential=Mock();event=threading.Event()
        with patch.object(batch.Azure,'call',side_effect=ValueError('transport interrupted')) as call:
            api=web.job_api(session,event)
            with self.assertRaises(ValueError):api.call('PUT',AID,{'properties':{}})
            another=web.job_api(session,event)
            with self.assertRaisesRegex(ValueError,'已提交'):another.call('PUT',AID,{'properties':{}})
        self.assertEqual(call.call_count,1)

    def test_web_defaults_to_maximum_and_can_select_tier1(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        self.assertEqual(web.configurations(session,data)[0][0]['quota_goal'],'maximum')
        self.assertEqual(web.configurations(session,{**data,'quota_goal':'tier1'})[0][0]['quota_goal'],'tier1')

    def test_empty_subscription_inventory_error_remains_visible_in_review(self):
        class Unreadable(MemoryAzure):
            def all(self,path,params=None):
                if path.endswith('/accounts'):raise ValueError('资源读取失败，不可判断不存在')
                return super().all(path,params)
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID,'displayName':'Azure Subscription A03'}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-01-01'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=Unreadable()):web.start_scan(session,data);session.thread.join(3)
        self.assertEqual(session.plans[0]['plan']['rows'][0]['status'],'blocked')
        self.assertEqual(session.plans[0]['plan']['rows'][0]['subscription_id'],SID)
        self.assertEqual(session.plans[0]['authoritative']['rows'],[])

    def test_actual_result_keeps_sku_and_real_request_url(self):
        row={'subscription':'A03','subscription_id':SID,'region':'eastus2','model':'gpt-image-2'}
        output={'account_id':AID,'model':{'name':'gpt-image-2','version':'actual'},'deployment':'image','capacity':2,'sku':'GlobalStandard','endpoints':{'AI Foundry API':'https://account.services.ai.azure.com/'}}
        actual=web.result_row(row,output,'project','existing')
        self.assertEqual(actual['sku'],'GlobalStandard')
        self.assertEqual(actual['url'],'https://account.services.ai.azure.com/openai/v1/images/generations')

    def test_customer_template_has_exact_headers_three_sheets_and_subscription_summary(self):
        from openpyxl import load_workbook
        rows=[{'订阅':'A03','订阅 ID':SID,'资源名称-foundry':'account','资源':'project','url':'https://account.services.ai.azure.com/openai/v1/images/generations','密钥':'synthetic','_account_id':AID,'status':'已核验','verified':True,'model':'gpt-image-2','sku':'GlobalStandard'}]
        book=load_workbook(io.BytesIO(web.excel(rows,rows,[{'subscriptionId':SID,'displayName':'A03'}])))
        self.assertEqual(book.sheetnames,['Sheet1','Sheet2','Sheet3'])
        self.assertEqual(book['Sheet1']['D1'].value,'资源名称-foundry项目')
        self.assertEqual(book['Sheet1']['F2'].value,'project')
        self.assertIsNone(book['Sheet1']['G2'].hyperlink)
        self.assertEqual(book['Sheet3'].max_row,2)

    def test_duplicate_existing_deployments_cannot_hide_an_unverified_target(self):
        from openpyxl import load_workbook
        rows=[{'订阅':'A03','订阅 ID':SID,'资源':'p','_account_id':AID,'model':'gpt-image-2','region_code':'eastus2','deployment':name,'verified':True} for name in ('one','two')]
        scopes=[{'subscriptionId':SID,'expected_models':2,'expected_targets':[['gpt-image-2','eastus2',AID.lower(),'one'],['gpt-image-2','westus3',AID.lower(),'three']]}]
        sheet=load_workbook(io.BytesIO(web.excel(rows,rows,scopes)))['Sheet3']
        values=dict(zip(next(sheet.values),list(sheet.values)[1]))
        self.assertEqual(values['结果'],'未完成核验')
        self.assertEqual(values['未完成项'],1)

    def test_missing_second_existing_deployment_is_not_hidden_in_export(self):
        from openpyxl import load_workbook
        session=web.Session()
        row={'subscription':'A03','subscription_id':SID,'region':'eastus2','model':'gpt-image-2','status':'existing','existing':[{'account_id':AID,'deployment':name} for name in ('one','two')]}
        session.plans=[{'plan':{'rows':[row]}}]
        session.results=[{'订阅':'A03','订阅 ID':SID,'资源':'p','_account_id':AID,'model':'gpt-image-2','region_code':'eastus2','deployment':'one','verified':True}]
        session.selected_subscriptions=[{'subscriptionId':SID,'expected_targets':[list(t) for t in web.required_targets(row)]}]
        details=web.export_details(session)
        self.assertEqual({r['deployment'] for r in details},{'one','two'})
        sheet=load_workbook(io.BytesIO(web.excel(session.results,details,session.selected_subscriptions)))['Sheet3']
        values=dict(zip(next(sheet.values),list(sheet.values)[1]))
        self.assertEqual(values['结果'],'未完成核验');self.assertEqual(values['未完成项'],1)

    def test_incompatible_subscription_does_not_stop_later_subscription(self):
        second='22222222-2222-4222-8222-222222222222';second_aid=AID.replace(SID,second)
        class Multiple(MemoryAzure):
            def all(self,path,params=None):
                if path=='/subscriptions':return [{'subscriptionId':s,'tenantId':TENANT,'state':'Enabled','displayName':'Azure Subscription '+str(i)} for i,s in enumerate([SID,second])]
                if path.endswith('/accounts'):
                    prefix=path.split('/providers/')[0]+'/'
                    return [v for k,v in self.resources.items() if k.startswith(prefix) and v.get('kind') in ('AIServices','OpenAI')]
                return super().all(path,params)
        api=Multiple()
        api.resources[AID]={'id':AID,'name':'cv-a03-001','kind':'OpenAI','location':'eastus2','properties':{'provisioningState':'Succeeded'}}
        api.resources[second_aid]={'id':second_aid,'name':'cv-a03-001','kind':'AIServices','location':'eastus2','identity':{'principalId':'synthetic'},'properties':{'provisioningState':'Succeeded','allowProjectManagement':True}}
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':s} for s in [SID,second]]
        data={'subscription_ids':[SID,second],'models':[{'name':'gpt-image-2','version':'2026-01-01'}],'regions':['eastus2'],'resource_group_template':'CV-{sub8}','account_template':'cv-{sub8}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(3)
            self.assertEqual([r['status'] for r in session.plans[0]['plan']['rows']],['blocked','planned'])
            try:web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True})
            except ValueError as error:self.fail('应继续可用订阅：'+str(error))
            session.thread.join(3)
        self.assertEqual(session.job['status'],'partial',session.job)
        self.assertTrue(any(p.startswith(second_aid+'/deployments/') for p in api.writes))
        self.assertTrue(all('/subscriptions/'+second+'/' in p for p in api.writes))
        self.assertEqual(session.results[0]['订阅 ID'],second)

    def test_legacy_reuse_does_not_validate_an_unused_project_name(self):
        aid=AID.rsplit('/',1)[0]+'/'+('a'*60)
        class Legacy(MemoryAzure):
            def all(self,path,params=None):
                if path.endswith('/accounts'):return [self.resources[aid]]
                return super().all(path,params)
        api=Legacy();api.resources[aid]={'id':aid,'name':'a'*60,'kind':'OpenAI','location':'eastus2','properties':{'provisioningState':'Succeeded'}}
        did=aid+'/deployments/image'
        api.resources[did]={'id':did,'name':'image','sku':{'name':'GlobalStandard','capacity':2},'properties':{'provisioningState':'Succeeded','model':{'name':'gpt-image-2','format':'OpenAI','version':'old'}}}
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-01-01'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):web.start_scan(session,data);session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(session.plans[0]['plan']['rows'][0]['status'],'existing')
        self.assertEqual(api.writes,[])

    def test_auto_regions_capacity_failure_still_allows_existing_model_inventory(self):
        class BrokenCapacity(MemoryAzure):
            def all(self,path,params=None):
                if path.endswith('/modelCapacities'):raise ValueError('invalid Azure capacity response')
                return super().all(path,params)
        api=BrokenCapacity();api.resources[AID]={'id':AID,'name':'cv-a03-001','kind':'AIServices','location':'eastus2','properties':{'provisioningState':'Succeeded'}}
        did=AID+'/deployments/text'
        api.resources[did]={'id':did,'name':'text','sku':{'name':'GlobalStandard','capacity':2},'properties':{'provisioningState':'Succeeded','model':{'name':'gpt-6-sol','format':'OpenAI','version':'old'}}}
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-6-sol','version':''}],'regions':[],'auto_regions':True,'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):web.start_scan(session,data);session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(session.plans[0]['plan']['rows'][0]['status'],'existing')
        self.assertEqual(api.writes,[])

    def test_existing_openai_deployment_is_reused_without_creating_project(self):
        class Legacy(MemoryAzure):
            def all(self,path,params=None):
                if path.endswith('/accounts'):return [self.resources[AID]]
                return super().all(path,params)
        api=Legacy()
        api.resources[AID]={'id':AID,'name':'cv-a03-001','kind':'OpenAI','location':'eastus2','properties':{'provisioningState':'Succeeded','endpoint':'https://example.invalid'}}
        did=AID+'/deployments/old-image'
        api.resources[did]={'id':did,'name':'old-image','sku':{'name':'DataZoneStandard','capacity':2},'properties':{'provisioningState':'Succeeded','model':{'name':'gpt-image-2','format':'OpenAI','version':'old'},'rateLimits':[{'key':'request','count':2,'renewalPeriod':60}]}}
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-01-01'}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(3)
            self.assertEqual(session.job['status'],'done',session.job)
            self.assertEqual(session.plans[0]['plan']['rows'][0]['status'],'existing')
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(api.writes,[])
        self.assertEqual(session.results[0]['deployment'],'old-image')
        self.assertEqual(session.results[0]['资源'],'')

    def test_existing_project_with_different_name_is_reused_without_writes(self):
        api=MemoryAzure()
        api.resources[AID]={'id':AID,'name':'cv-a03-001','kind':'AIServices','location':'eastus2','properties':{'provisioningState':'Succeeded','allowProjectManagement':True}}
        did=AID+'/deployments/customer-image'
        api.resources[did]={'id':did,'name':'customer-image','sku':{'name':'GlobalStandard','capacity':2},'properties':{'provisioningState':'Succeeded','model':{'name':'gpt-image-2','format':'OpenAI','version':'old'}}}
        pid=AID+'/projects/customer-project'
        api.resources[pid]={'id':pid,'name':'customer-project','properties':{'provisioningState':'Succeeded'}}
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':''}],'regions':['eastus2'],'resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(3)
            self.assertEqual(session.plans[0]['projects'][AID]['name'],'customer-project')
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(session.results[0]['资源'],'customer-project')
        self.assertEqual(api.writes,[])

    def test_default_quota_mode_does_not_require_manual_scope_confirmation(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-04-21'}],'regions':['eastus2'],'resource_group_template':'rg-{sub8}','account_template':'ai-{sub8}-{region}','project_template':'{foundry}-project'}
        configs,_=web.configurations(session,data);self.assertEqual(configs[0]['quota_scope'],'auto')

    def test_manual_other_regions_are_separate_from_image2_candidates(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-04-21'},{'name':'gpt-6-sol','version':'2026-01-01'}],'regions':list(web.REGIONS),'other_regions':['eastus2'],'auto_regions':False,'resource_group_template':'rg-{sub8}','account_template':'ai-{sub8}-{region}','project_template':'{foundry}-project'}
        configs,_=web.configurations(session,data)
        self.assertEqual(configs[0]['regions'],list(web.REGIONS));self.assertEqual(configs[1]['regions'],['eastus2'])

    def test_only_text_models_do_not_require_image2_region_selection(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-6-sol','version':'2026-01-01'}],'regions':[],'other_regions':['eastus2'],'auto_regions':False,'resource_group_template':'rg-{sub8}','account_template':'ai-{sub8}-{region}','project_template':'{foundry}-project'}
        configs,_=web.configurations(session,data);self.assertEqual(configs[0]['regions'],['eastus2'])

    def test_cancel_invalidates_pending_browser_authorization(self):
        session=web.Session();session.login_generation=2;session.oauth={'flow':{'state':'synthetic-state'}};session.login_url='https://login.microsoftonline.com/'
        web.cancel_job(session)
        self.assertIsNone(session.oauth);self.assertEqual(session.login_url,'');self.assertEqual(session.login_generation,3)

    def test_browser_callback_reads_subscriptions_without_resource_writes(self):
        session=web.Session();session.login_generation=1;client=Mock();client.acquire_token_by_auth_code_flow.return_value={'access_token':'synthetic-token'}
        session.oauth={'client':client,'flow':{'state':'synthetic-state'},'expires':time.monotonic()+60,'expected_tenant':TENANT,'generation':1}
        credential=Mock();credential.tenant=TENANT;api=MemoryAzure()
        with patch.object(web,'BrowserCredential',return_value=credential),patch.object(web,'JobAzure',return_value=api):
            web.finish_browser_login(session,{'state':'synthetic-state','code':'synthetic-code'});session.thread.join(3)
        self.assertEqual(session.job['status'],'done');self.assertEqual(session.subscriptions[0]['subscriptionId'],SID)
        self.assertIs(session.credential,credential);self.assertEqual(api.writes,[])

    def test_logout_during_browser_redemption_cannot_restore_connection(self):
        session=web.Session();client=Mock();session.oauth={'client':client,'flow':{'state':'synthetic-state'},'expires':time.monotonic()+60,'generation':0}
        def redeem(*args):session.login_generation+=1;return {'access_token':'synthetic-token'}
        client.acquire_token_by_auth_code_flow.side_effect=redeem
        with patch.object(web,'BrowserCredential',return_value=Mock()):
            with self.assertRaises(ValueError):web.finish_browser_login(session,{'state':'synthetic-state','code':'synthetic-code'})
        self.assertIsNone(session.api);self.assertIsNone(session.credential)

    def test_browser_login_uses_original_arm_delegated_scope(self):
        session=web.Session();client=Mock();client.initiate_auth_code_flow.return_value={'state':'synthetic-state','auth_uri':'https://login.microsoftonline.com/organizations/oauth2/v2.0/authorize'}
        with patch.object(web,'oauth_client',return_value=client),patch.dict('os.environ',{'OAUTH_CLIENT_ID':SID,'OAUTH_CLIENT_SECRET':'synthetic-secret'}):
            web.start_login(session,{'tenant_id':''});session.thread.join(3)
        self.assertEqual(client.initiate_auth_code_flow.call_args.args[0],['https://management.azure.com/user_impersonation'])
        self.assertTrue(session.login_url.startswith('https://login.microsoftonline.com/'))
        self.assertNotIn('synthetic-secret',str(session.snapshot()))

    def test_browser_callback_is_browser_bound_and_single_use(self):
        first=web.Session();second=web.Session();client=Mock()
        first.oauth={'client':client,'flow':{'state':'synthetic-state'},'expires':time.monotonic()+60,'expected_tenant':None}
        with self.assertRaises(ValueError):web.finish_browser_login(second,{'state':'synthetic-state','code':'synthetic-code'})
        with self.assertRaises(ValueError):web.finish_browser_login(first,{'state':'different','code':'synthetic-code'})
        client.acquire_token_by_auth_code_flow.return_value={'error':'access_denied','error_description':'secret=never-echo'}
        with self.assertRaises(ValueError) as failure:web.finish_browser_login(first,{'state':'synthetic-state','code':'synthetic-code'})
        self.assertNotIn('never-echo',str(failure.exception));self.assertIsNone(first.oauth)
        with self.assertRaises(ValueError):web.finish_browser_login(first,{'state':'synthetic-state','code':'synthetic-code'})
        self.assertEqual(client.acquire_token_by_auth_code_flow.call_count,1)

    def test_logout_between_login_check_and_commit_cannot_restore_credential(self):
        session=web.Session();gate={'commit':False};original=session.lock
        class LogoutAtCommit:
            def __enter__(self):
                original.acquire()
                if gate['commit']:
                    gate['commit']=False;session.cancel.set();session.api=None;session.credential=None
                return self
            def __exit__(self,*args):original.release()
        session.lock=LogoutAtCommit()
        def subs(*args):gate['commit']=True;return [{'subscriptionId':SID,'displayName':'Synthetic'}]
        with patch('azure.identity.DeviceCodeCredential',return_value=Mock()),patch.object(batch,'subscriptions',side_effect=subs):
            web.start_login(session,{'tenant_id':TENANT});session.thread.join(3)
        self.assertEqual(session.job['status'],'cancelled');self.assertIsNone(session.api);self.assertIsNone(session.credential)

    def test_empty_version_resolves_azure_default_for_each_region(self):
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID,'displayName':'Azure Subscription A03'}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':''}],'regions':['eastus2','westus3'],'quota_scope':'regional','resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=MemoryAzure()):web.start_scan(session,data);session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        rows=session.plans[0]['plan']['rows'];self.assertEqual({r['version'] for r in rows},{'2026-01-01'})
        self.assertTrue(all(r['capacity']==1 for r in rows));self.assertEqual(len(rows),2)

    def test_default_version_requires_explicit_azure_default_marker(self):
        rows=[{'model':{'name':'gpt-6-sol','format':'OpenAI','version':'v2','isDefaultVersion':True,'skus':[{'name':'GlobalStandard'}]}},{'model':{'name':'gpt-6-sol','format':'OpenAI','version':'v9','skus':[{'name':'GlobalStandard'}]}}]
        self.assertEqual(web.default_version(rows,'gpt-6-sol','GlobalStandard'),'v2')
        with self.assertRaises(ValueError):web.default_version(rows[1:],'gpt-6-sol','GlobalStandard')

    def test_cancel_interrupts_rate_limit_wait(self):
        event=threading.Event();event.set();api=web.JobAzure(Mock(),event,lambda _:None)
        with self.assertRaises(web.Cancelled):api.pause(120)

    def test_two_models_reuse_accounts_and_projects_and_region_numbers(self):
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT
        session.subscriptions=[{'subscriptionId':SID,'displayName':'Azure Subscription A03'}]
        api=MemoryAzure()
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-01-01'},{'name':'gpt-6-sol','version':'2026-01-01'}],'regions':['eastus2','westus3'],'auto_regions':True,'quota_scope':'regional','resource_group_template':'CV-{subcode}','account_template':'cv-{subcode}-{index:03d}','project_template':'{foundry}-resource'}
        with patch.object(web,'JobAzure',return_value=api):
            web.start_scan(session,data);session.thread.join(3)
            self.assertEqual(session.job['status'],'done',session.job)
            web.start_deploy(session,{'plan_id':session.plan_id,'confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'done',session.job)
        self.assertEqual(len(session.results),5)
        accounts=[p for p in api.writes if '/accounts/' in p and '/deployments/' not in p and '/projects/' not in p]
        projects=[p for p in api.writes if '/projects/' in p]
        self.assertEqual(len(accounts),3);self.assertEqual(len(projects),3)
        self.assertEqual(len(set(accounts)),3)
        names={r['区域']:r['资源名称-foundry'] for r in session.results}
        self.assertEqual(names['East US 2'],'cv-a03-001');self.assertEqual(names['West US 3'],'cv-a03-002');self.assertEqual(names['eastus'],'cv-a03-006')

    def deployment_session(self):
        session=web.Session();session.api=Mock();session.credential=Mock();session.tenant=TENANT;session.plan_id='reviewed'
        row={'subscription':'Azure Subscription A03','subscription_id':SID,'region':'eastus2','model':'gpt-image-2','version':'requested','status':'existing','sku':'GlobalStandard'}
        plan={'rows':[row],'errors':[],'config':{},'tenant_id':TENANT}
        session.plans=[{'plan':plan,'projects':{AID:{'name':'project','new':False}}}]
        output={'account_id':AID,'model':{'name':'gpt-image-2','version':'actual-existing'},'deployment':'existing','capacity':2,'endpoint':'https://example.invalid'}
        return session,{'rows':[{**row,'outputs':[output]}],'errors':[]}

    def test_existing_deployment_exports_actual_version(self):
        session,report=self.deployment_session()
        with patch.object(batch,'execute',return_value=report),patch.object(web,'ensure_project'):
            web.start_deploy(session,{'plan_id':'reviewed','confirmed':True});session.thread.join(3)
        self.assertEqual(session.results[0]['version'],'actual-existing')

    def test_cancel_preserves_completed_resource_records(self):
        session,report=self.deployment_session()
        def execute(api,plan,on_progress):on_progress(report);raise web.Cancelled()
        with patch.object(batch,'execute',side_effect=execute):
            web.start_deploy(session,{'plan_id':'reviewed','confirmed':True});session.thread.join(3)
        self.assertEqual(session.job['status'],'cancelled')
        self.assertEqual(len(session.results),1)
        self.assertEqual(session.results[0]['version'],'actual-existing')
        self.assertEqual(session.results[0]['资源'],'')

    def test_image2_restriction_does_not_limit_other_models(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        data={'subscription_ids':[SID],'models':[{'name':'gpt-image-2','version':'2026-04-21'},{'name':'gpt-6-sol','version':'2026-01-01'}],'regions':['eastus2','eastus'],'auto_regions':True,'resource_group_template':'rg-{sub8}','account_template':'ai-{sub8}-{region}','project_template':'{foundry}-project'}
        configs,_=web.configurations(session,data)
        self.assertEqual(configs[0]['regions'],['eastus2']);self.assertEqual(configs[1]['regions'],'auto')

    def test_excel_has_exact_customer_columns_and_text_values(self):
        from openpyxl import load_workbook
        rows=[{'订阅':'=unsafe','订阅 ID':SID,'资源组':'CV-A03','资源名称-foundry':'cv-a03-001','区域':'East US 2','资源':'cv-a03-001-resource','url':'https://example.invalid','密钥':'synthetic-key'}]
        book=load_workbook(io.BytesIO(web.excel(rows)))
        self.assertEqual(list(next(book.active.values)),list(web.TEMPLATE_COLUMNS))
        self.assertEqual(book.active['A2'].data_type,'s')
        self.assertEqual(book.active['H2'].value,'synthetic-key')

    def test_review_cannot_be_executed_using_browser_supplied_plan(self):
        session=web.Session();session.api=Mock();session.tenant=TENANT
        with self.assertRaises(ValueError):web.start_deploy(session,{'plan_id':'made-up','confirmed':True})
        session.api.put.assert_not_called()

    def test_session_snapshot_does_not_expose_keys_or_credentials(self):
        session=web.Session();session.credential=Mock();session.keys={AID:'synthetic-key'}
        session.results=[{'密钥':'synthetic-key','订阅 ID':SID}]
        snapshot=json.dumps(session.snapshot())
        self.assertNotIn('synthetic-key',snapshot)
        self.assertNotIn('credential',snapshot)

    def test_cancel_stops_next_azure_call(self):
        event=threading.Event();event.set()
        api=web.JobAzure(Mock(),event,lambda _:None)
        with patch.object(batch.Azure,'call') as request:
            with self.assertRaises(web.Cancelled):api.call('PUT',AID,{})
        request.assert_not_called()

    def test_existing_project_is_reused_without_write(self):
        api=Mock();api.optional.return_value={'id':AID+'/projects/p','location':'eastus2','properties':{'provisioningState':'Succeeded'}}
        web.ensure_project(api,AID,'p','eastus2')
        api.put.assert_not_called()

    def test_missing_project_requires_parent_support_before_write(self):
        api=Mock();api.optional.return_value=None;api.get.return_value={'kind':'OpenAI','properties':{}}
        with self.assertRaises(ValueError):web.ensure_project(api,AID,'p','eastus2')
        api.put.assert_not_called()

    def test_config_only_accepts_visible_selected_subscriptions(self):
        session=web.Session();session.tenant=TENANT;session.subscriptions=[{'subscriptionId':SID}]
        with self.assertRaises(ValueError):web.configurations(session,{'subscription_ids':['44444444-4444-4444-8444-444444444444']})

    def test_device_login_challenge_is_returned_without_token(self):
        session=web.Session()
        def factory(**kwargs):
            kwargs['prompt_callback']('https://microsoft.com/devicelogin','TEST-CODE',None)
            result=Mock();result.get_token.side_effect=RuntimeError('token=do-not-leak');return result
        with patch('azure.identity.DeviceCodeCredential',side_effect=factory):
            web.start_login(session,{'tenant_id':TENANT});session.thread.join(3)
        self.assertNotIn('do-not-leak',json.dumps(session.snapshot()))
        self.assertEqual(session.snapshot()['job']['status'],'failed')


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=web.make_server('127.0.0.1',0)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.origin=f'http://127.0.0.1:{cls.server.server_port}'

    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close()

    def bootstrap(self):
        with urlopen(self.origin+'/') as response:
            self.assertIn('Foundry',response.read().decode());return response.headers['Set-Cookie'].split(';')[0]

    def test_session_isolation_and_no_cache(self):
        first=self.bootstrap();second=self.bootstrap();self.assertNotEqual(first,second)
        with urlopen(Request(self.origin+'/api/state',headers={'Cookie':first})) as response:
            self.assertEqual(response.headers['Cache-Control'],'no-store')
            self.assertFalse(json.load(response)['connected'])

    def test_invalid_callback_cannot_release_running_job_guard(self):
        cookie=self.bootstrap();session=self.server.sessions[cookie.split('=',1)[1]]
        session.job={'status':'running','title':'deployment'}
        with urlopen(Request(self.origin+'/auth/callback?state=invalid',headers={'Cookie':cookie})) as response:response.read()
        self.assertEqual(session.job['status'],'running')
        with self.assertRaises(ValueError):web.start(session,'second',lambda event:None)

    def test_unused_anonymous_sessions_cannot_permanently_fill_capacity(self):
        with self.server.sessions_lock:
            self.server.sessions.clear()
            for i in range(256):self.server.sessions['unused-'+str(i)]=web.Session()
        self.bootstrap()

    def test_uncompleted_browser_logins_cannot_permanently_fill_capacity(self):
        with self.server.sessions_lock:
            self.server.sessions.clear()
            for i in range(256):
                session=web.Session();session.oauth={'expires':time.monotonic()+60 if i%2 else time.monotonic()-1};session.job={'status':'done'}
                self.server.sessions['pending-'+str(i)]=session
        self.bootstrap()

    def test_cross_origin_write_is_rejected_before_body_read(self):
        cookie=self.bootstrap()
        request=Request(self.origin+'/api/login',data=b'{}',headers={'Cookie':cookie,'Origin':'https://evil.invalid','Content-Type':'application/json','X-Requested-With':'FoundryUI'})
        with self.assertRaises(HTTPError) as error:urlopen(request)
        self.assertEqual(error.exception.code,403)

    def test_key_export_requires_explicit_post(self):
        cookie=self.bootstrap()
        with self.assertRaises(HTTPError) as error:urlopen(Request(self.origin+'/api/export',headers={'Cookie':cookie}))
        self.assertEqual(error.exception.code,404)

    def test_unconnected_client_cannot_scan(self):
        cookie=self.bootstrap()
        request=Request(self.origin+'/api/scan',data=b'{}',headers={'Cookie':cookie,'Origin':self.origin,'Content-Type':'application/json','X-Requested-With':'FoundryUI'})
        with self.assertRaises(HTTPError) as error:urlopen(request)
        self.assertEqual(error.exception.code,400)

    def test_public_https_origin_uses_secure_cookie_and_rejects_http_origin(self):
        with patch.object(web,'PUBLIC_ORIGIN','https://workspace.example.invalid'):
            with urlopen(Request(self.origin+'/',headers={'Host':'workspace.example.invalid'})) as response:
                cookie=response.headers['Set-Cookie'];self.assertIn('; Secure',cookie)
                self.assertIn('SameSite=Lax',cookie);self.assertIn('HttpOnly',cookie)
            request=Request(self.origin+'/api/login',data=b'{}',headers={'Host':'workspace.example.invalid','Cookie':cookie.split(';')[0],'Origin':'http://workspace.example.invalid','Content-Type':'application/json','X-Requested-With':'FoundryUI'})
            with self.assertRaises(HTTPError) as error:urlopen(request)
            self.assertEqual(error.exception.code,403)


if __name__=='__main__':unittest.main()
