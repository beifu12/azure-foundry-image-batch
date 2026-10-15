#!/usr/bin/env python3
"""Customer-owned, loopback-only browser workspace. Tokens and keys stay in RAM."""
import argparse
import copy
import io
import json
import logging
import os
import re
import secrets
import threading
import time
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit,parse_qs
from uuid import UUID
from concurrent.futures import ThreadPoolExecutor, as_completed

import image_batch as batch

ROOT=Path(__file__).resolve().parent
COLUMNS=('订阅','订阅 ID','资源组','资源名称-foundry','区域','资源','url','密钥')
TEMPLATE_COLUMNS=(*COLUMNS[:3],'资源名称-foundry项目',*COLUMNS[4:])
MAX_SUBSCRIPTIONS_IN_FLIGHT=3
MAX_CHECK_SUBSCRIPTIONS=8
MAX_CHECK_READS=16
CHECK_READ_SLOTS=threading.BoundedSemaphore(MAX_CHECK_READS)
WEB_VERSION='2026.10.15.4'
REGIONS={'eastus2':'East US 2','westus3':'West US 3','polandcentral':'Poland Central','swedencentral':'Sweden Central','uaenorth':'UAE North'}
ACTIVE={'running','cancelling'}
PUBLIC_ORIGIN=os.environ.get('PUBLIC_ORIGIN','').rstrip('/')
WRITE_LOCK=threading.Lock()
LOGIN_SLOTS=threading.BoundedSemaphore(8)
ARM_SCOPES=['https://management.azure.com/user_impersonation']


class Cancelled(Exception):pass


class Session:
    def __init__(self):
        self.lock=threading.RLock();self.cancel=threading.Event();self.thread=None
        self.tenant='';self.api=None;self.credential=None;self.subscriptions=[]
        self.job={};self.challenge=None;self.plans=[];self.plan_id='';self.results=[];self.keys={};self.catalog={};self.issues=[]
        self.touched=time.monotonic()
        self.oauth=None;self.login_url='';self.login_generation=0;self.login_error=''
        self.selected_subscriptions=[];self.submitted_puts=set()

    def snapshot(self):
        with self.lock:
            return copy.deepcopy({'connected':self.api is not None,'tenant_id':self.tenant,'subscriptions':self.subscriptions,
                'job':self.job,'challenge':self.challenge,'plans':[{'plan':{'rows':b['plan']['rows'],'errors':b['plan']['errors']},'projects':b['projects']} for b in self.plans],'plan_id':self.plan_id,
                'login_url':self.login_url,'login_error':self.login_error,'login_mode':'browser' if os.environ.get('OAUTH_CLIENT_ID') else 'device',
                'results':[{**r,'密钥':'••••••••' if self.keys.get(r.get('_account_id')) else '未读取'} for r in self.results],
                'catalog':self.catalog,'issues':self.issues,'models':batch.MODELS})


class JobAzure(batch.Azure):
    def __init__(self,credential,event,notify):
        super().__init__(credential);self.event=event;self.notify=notify;self.quiet=True

    def call(self,method,path,body=None,params=None,headers=None):
        if self.event.is_set():raise Cancelled()
        cache=getattr(self,'read_cache',None)
        if cache is not None and method!='GET':raise ValueError('检查库存为只读阶段，不允许提交创建或密钥请求。')
        cache_key=(path,json.dumps(params,sort_keys=True))
        if method=='GET' and cache is not None and cache_key in cache:return copy.deepcopy(cache[cache_key])
        failures=getattr(self,'read_errors',{})
        if cache is not None and cache_key in failures:raise failures[cache_key]
        if method=='PUT' and hasattr(self,'submitted_puts'):
            with self.guard_lock:
                key=urlsplit(path).path.lower()
                if key in self.submitted_puts:raise ValueError('此资源已提交过创建请求；需确认 Azure 实际状态，不自动重放写入。')
                self.submitted_puts.add(key)
        self.read_deadline=time.monotonic()+180
        tail=urlsplit(path).path.rsplit('/',1)[-1]
        label={'models':'读取模型版本','modelCapacities':'读取可用区域','usages':'核对配额','accounts':'检查 Foundry','deployments':'检查现有部署','listKeys':'读取所选资源密钥','subscriptions':'读取启用订阅'}.get(tail,'读取资源状态' if method=='GET' else '提交 Azure 配置')
        self.notify(label)
        if method=='PUT' and '/accounts/' in path and '/deployments/' not in path and '/projects/' not in path:
            body=copy.deepcopy(body);body['identity']={'type':'SystemAssigned'};body['properties']['allowProjectManagement']=True
        if cache is not None:
            while not CHECK_READ_SLOTS.acquire(timeout=.1):
                if self.event.is_set():raise Cancelled()
            try:
                if self.event.is_set():raise Cancelled()
                result=super().call(method,path,body,params,headers)
            except (ValueError,RuntimeError) as error:
                self.read_errors[cache_key]=error
                raise
            finally:CHECK_READ_SLOTS.release()
        else:result=super().call(method,path,body,params,headers)
        if self.event.is_set():raise Cancelled()
        if method=='GET' and cache is not None:cache[cache_key]=copy.deepcopy(result)
        return result

    def pause(self,seconds):
        if time.monotonic()+seconds>getattr(self,'read_deadline',float('inf')):raise RuntimeError('读取限流等待已超过三分钟预算；已停止重试，请稍后重新检查。')
        self.notify(f'Azure 限流或操作仍在配置，等待 {seconds} 秒；可取消后续步骤。')
        if self.event.wait(seconds):raise Cancelled()


def progress(session,completed=None,total=None,message=None,phase=None):
    with session.lock:
        if completed is not None:session.job['completed']=completed
        if total is not None:session.job['total']=total
        if message is not None:session.job['message']=message
        if phase is not None:session.job['phase']=phase


def job_api(session,event,notify=None,snapshot=False):
    api=JobAzure(session.credential,event,notify or (lambda message:progress(session,message=message)))
    api.submitted_puts=session.submitted_puts;api.guard_lock=session.lock
    if snapshot:api.read_cache={};api.read_errors={}
    return api


def prefetch_reads(api,requests):
    """Warm this check's read snapshot with independent HTTP sessions, never writes."""
    if not isinstance(api,batch.Azure):return
    unique={(path,json.dumps(params,sort_keys=True)):(path,params) for path,params in requests}
    workers=threading.local();transports=[]
    def initialize():
        local=type(api)(api.credential,api.event,api.notify)
        local.read_cache=api.read_cache;local.read_errors=api.read_errors
        local.session.trust_env=api.session.trust_env
        workers.api=local;transports.append(local.session)
    def read(item):
        if api.event.is_set():raise Cancelled()
        try:workers.api.all(*item)
        except (ValueError,RuntimeError):pass  # Retained as a failure, never an empty list.
    try:
        with ThreadPoolExecutor(max_workers=4,initializer=initialize) as pool:
            futures=[pool.submit(read,item) for item in unique.values()]
            for future in as_completed(futures):future.result()
    finally:
        for transport in transports:transport.close()


def start(session,title,operation):
    with session.lock:
        if session.job.get('status') in ACTIVE:raise ValueError('当前任务仍在执行；可先取消，待正在处理的请求返回后重试。')
        session.cancel=threading.Event();event=session.cancel
        session.job={'title':title,'status':'running','completed':0,'total':0,'message':'正在准备','phase':'准备','started_at':time.time()}
        before_operation=getattr(session,'before_operation',None)
        def run():
            try:
                if before_operation:before_operation()
                operation(event)
                if event.is_set():raise Cancelled()
                with session.lock:
                    partial=session.job.get('partial',False)
                    session.job.update(status='partial' if partial else 'done',message='已结束，部分项目需处理；请查看结果与提醒。' if partial else '任务完成',completed=session.job.get('total',0))
            except Cancelled:
                with session.lock:session.job.update(status='cancelled',message='已停止后续步骤；已提交到 Azure 的操作可能仍在执行，请先重新检查。')
            except (batch.AzureError,ValueError) as error:
                with session.lock:session.job.update(status='failed',message=str(error))
            except Exception:
                with session.lock:session.job.update(status='failed',message='请求未完成。请检查登录、网络及组织策略，再重新检查；不会自动重放创建或密钥请求。')
            finally:
                with session.lock:session.challenge=None
        session.thread=threading.Thread(target=run,daemon=True);session.thread.start()


def oauth_client(tenant):
    import msal
    client_id=os.environ.get('OAUTH_CLIENT_ID');secret=os.environ.get('OAUTH_CLIENT_SECRET')
    if not client_id or not secret:raise ValueError('服务端 Microsoft 浏览器登录尚未配置。')
    return msal.ConfidentialClientApplication(str(UUID(client_id)),client_credential=secret,authority='https://login.microsoftonline.com/'+tenant,token_cache=msal.TokenCache(),timeout=30)


def cancel_job(session):
    with session.lock:
        session.cancel.set();session.oauth=None;session.login_url='';session.login_generation+=1
        if session.job.get('status') in ACTIVE:
            progress(session,message='正在停止后续步骤，等待当前 Azure 请求返回。');session.job['status']='cancelling'


class BrowserCredential:
    def __init__(self,client,result):
        from azure.core.credentials import AccessToken
        claims=result.get('id_token_claims') or {}
        self.tenant=str(UUID(claims.get('tid','')));user=str(UUID(claims.get('oid','')))
        accounts=[a for a in client.get_accounts() if a.get('local_account_id')==user and a.get('realm') in (self.tenant,'organizations')]
        if len(accounts)!=1:raise ValueError('无法唯一确认客户登录身份，请重新连接。')
        self.client=client;self.account=accounts[0];self.lock=threading.Lock()
        self.name=claims.get('preferred_username') or claims.get('name') or ''
        self.token_type=AccessToken

    def get_token(self,*scopes,**kwargs):
        if scopes not in ((batch.ARM+'/.default',),tuple(ARM_SCOPES)):raise ValueError('此客户连接仅用于 Azure 管理接口。')
        with self.lock:result=self.client.acquire_token_silent(ARM_SCOPES,account=self.account)
        if not result or not result.get('access_token'):raise ValueError('客户登录已过期或需要补充授权，请重新连接。')
        return self.token_type(result['access_token'],int(result.get('expires_on',time.time()+int(result.get('expires_in',300)))))


def start_login(session,data):
    if not os.environ.get('OAUTH_CLIENT_ID'):return start_device_login(session,data)
    tenant=str(UUID(data['tenant_id'])) if data.get('tenant_id') else 'organizations'
    if session.api is not None:raise ValueError('请先退出当前客户连接。')
    with session.lock:
        if session.job.get('status') in ACTIVE:raise ValueError('当前连接步骤仍在执行，请稍后重试。')
        session.login_generation+=1;generation=session.login_generation;session.login_error=''
    def begin(event):
        if not LOGIN_SLOTS.acquire(blocking=False):raise ValueError('当前登录请求较多，请稍后重试。')
        try:
            client=oauth_client(tenant)
            flow=client.initiate_auth_code_flow(ARM_SCOPES,redirect_uri=(PUBLIC_ORIGIN or 'http://127.0.0.1:8767')+'/auth/callback',prompt='select_account',response_mode='query')
        finally:LOGIN_SLOTS.release()
        if not flow.get('state') or urlsplit(flow.get('auth_uri','')).hostname!='login.microsoftonline.com':raise ValueError('无法开始 Microsoft 浏览器登录。')
        with session.lock:
            if event.is_set() or session.login_generation!=generation:raise Cancelled()
            session.oauth={'client':client,'flow':flow,'expires':time.monotonic()+900,'expected_tenant':None if tenant=='organizations' else tenant,'generation':generation}
            session.login_url=flow['auth_uri']
    start(session,'跳转 Microsoft 选择 CSP / 客户账号',begin)


def finish_browser_login(session,response):
    with session.lock:
        pending=session.oauth
        if not pending or pending['expires']<time.monotonic() or not response.get('state') or not secrets.compare_digest(str(pending['flow']['state']),response['state']):raise ValueError('授权已过期、已使用或不属于当前浏览器，请重新登录。')
        session.oauth=None;session.login_url=''
    # Redeem once; neither the authorization code nor SDK error text reaches the UI.
    try:result=pending['client'].acquire_token_by_auth_code_flow(pending['flow'],response)
    except Exception:raise ValueError('Microsoft 客户授权未完成，请重新连接。') from None
    if not result.get('access_token'):raise ValueError('Microsoft 客户授权未完成；请按 Microsoft 提示授予 Azure 委托访问权限。')
    try:credential=BrowserCredential(pending['client'],result)
    except Exception:raise ValueError('无法确认客户登录身份，请重新连接。') from None
    if pending.get('expected_tenant') and credential.tenant!=pending['expected_tenant']:raise ValueError('登录账号与指定的客户租户不一致。')
    def connect(event):
        api=JobAzure(credential,event,lambda message:progress(session,message=message))
        subs=batch.subscriptions(api,{'tenant_id':credential.tenant,'subscription_ids':'all'})
        with session.lock:
            if event.is_set() or session.login_generation!=pending.get('generation',0):raise Cancelled()
            session.api=batch.Azure(credential);session.credential=credential;session.tenant=credential.tenant
            session.subscriptions=sorted(subs,key=lambda r:[int(s) if s.isdigit() else s.lower() for s in re.split(r'(\d+)',r.get('displayName',''))])
    with session.lock:
        if session.login_generation!=pending.get('generation',0):raise ValueError('当前客户登录已退出或被新连接替换。')
        start(session,'连接客户账号并读取可访问订阅',connect)


def start_device_login(session,data):
    tenant=str(UUID(data.get('tenant_id','')))
    def login(event):
        if not LOGIN_SLOTS.acquire(blocking=False):raise ValueError('当前登录请求较多，请稍后重试。')
        try:login_locked(event)
        finally:LOGIN_SLOTS.release()
    def login_locked(event):
        from azure.identity import DeviceCodeCredential
        def prompt(uri,code,expires):
            if urlsplit(uri).hostname not in ('microsoft.com','www.microsoft.com','login.microsoft.com','login.microsoftonline.com'):raise ValueError('Microsoft 登录地址不符合预期。')
            with session.lock:session.challenge={'url':uri,'code':code}
            progress(session,message='在 Microsoft 页面完成登录；本工具不收集密码和验证码。')
        credential=DeviceCodeCredential(tenant_id=tenant,prompt_callback=prompt,timeout=300,authority='https://login.microsoftonline.com')
        credential.get_token(batch.ARM+'/.default')
        api=JobAzure(credential,event,lambda message:progress(session,message=message))
        subs=batch.subscriptions(api,{'tenant_id':tenant,'subscription_ids':'all'})
        if event.is_set():raise Cancelled()
        with session.lock:
            if event.is_set():raise Cancelled()
            session.credential=credential;session.api=batch.Azure(credential);session.tenant=tenant
            session.subscriptions=sorted(subs,key=lambda r:[int(s) if s.isdigit() else s.lower() for s in re.split(r'(\d+)',r.get('displayName',''))])
            session.plans=[];session.plan_id='';session.results=[];session.keys={};session.catalog={};session.issues=[]
    if session.api is not None:raise ValueError('切换账号前请先退出当前连接。')
    start(session,'登录并读取订阅',login)


def configurations(session,data):
    selected=data.get('subscription_ids')
    if not isinstance(selected,list) or not selected:raise ValueError('请至少勾选一个订阅。')
    visible={s['subscriptionId'] for s in session.subscriptions}
    selected=list(dict.fromkeys(str(UUID(s)) for s in selected))
    if not set(selected)<=visible:raise ValueError('所选订阅不在当前登录账号可访问的范围内。')
    models=data.get('models',[])
    if not isinstance(models,list) or not models or len(models)>len(batch.MODELS):raise ValueError('请至少选择一个支持的模型，并填写版本。')
    if len({r.get('name') for r in models})!=len(models):raise ValueError('模型选择不能重复。')
    regions=data.get('regions')
    others=data.get('other_regions',['eastus2'])
    for values in (regions,others):
        if not isinstance(values,list) or len(values)>30 or any(not isinstance(r,str) or not re.fullmatch(r'[a-z0-9]{2,40}',r) for r in values):raise ValueError('区域须为最多 30 个有效 Azure 区域代码。')
    result=[]
    for model in models:
        approved=list(dict.fromkeys(r for r in regions if r in REGIONS)) if model['name']=='gpt-image-2' else ('auto' if data.get('auto_regions') is True else list(dict.fromkeys(others)))
        if not approved:raise ValueError('请为所选模型指定至少一个部署区域。')
        version=model.get('version','').strip()
        config=dict(tenant_id=session.tenant,subscription_ids=selected,model=model['name'],model_version=version or 'AzureDefault',use_default=not version,sku=data.get('sku','GlobalStandard'),regions=approved,quota_goal=data.get('quota_goal','maximum'),quota_scope=data.get('quota_scope','auto'),data_zones=data.get('data_zones',{}),resource_group_template=data.get('resource_group_template',''),account_template=data.get('account_template',''),deployment_name=model['name'])
        batch.check_config(config)
        result.append(config)
    project=data.get('project_template','')
    if not isinstance(project,str) or not project:raise ValueError('请填写 Project 命名规则。')
    return result,project


def project_name(template,row,account):
    # {foundry} always resolves to the actual reused or new resource name.
    template=template.replace('{foundry}',account.rsplit('/',1)[-1])
    return batch.render_name(template,row['subscription_id'],row['region'],'project',row['subscription'],row.get('region_index',1))


def default_version(rows,name,sku):
    matches=[]
    for raw in rows:
        model=raw.get('model') or raw
        if model.get('format')=='OpenAI' and model.get('name')==name and model.get('isDefaultVersion') is True and any(s.get('name')==sku for s in model.get('skus',[])):
            matches.append(model.get('version'))
    versions=set(matches)
    if len(versions)!=1 or not next(iter(versions),None):raise ValueError('Azure 未返回唯一默认版本；请在高级设置中填写精确版本。')
    return next(iter(versions))


def ensure_project(api,account,name,region):
    path=account+'/projects/'+name
    existing=api.optional(path)
    if existing is not None:
        if existing.get('properties',{}).get('provisioningState')!='Succeeded':raise ValueError('同名 Project 尚未成功；不会覆盖或重复创建。')
        return existing
    parent=api.get(account)
    if parent.get('kind')!='AIServices' or not parent.get('properties',{}).get('allowProjectManagement') or not parent.get('identity',{}).get('principalId'):
        raise ValueError('现有资源不支持 Project 管理或没有托管身份；不会自动升级/覆盖它。')
    api.put(path,{'location':region,'identity':{'type':'SystemAssigned'},'properties':{}})
    return batch.wait_ready(api,path)


def scan_scope(session,configs,template,event,notify,subscriptions_cache=None):
    api=job_api(session,event,lambda message:notify(message=message),snapshot=True)
    if subscriptions_cache is not None:api.read_cache.update(copy.deepcopy(subscriptions_cache))
    initial=[]
    for sid in configs[0]['subscription_ids']:
        root=f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices'
        initial.append((root+'/accounts',None))
        regions={region for config in configs if config['regions']!='auto' for region in config['regions']}
        for region in sorted(regions):
            initial.extend([(root+f'/locations/{region}/models',None),
                            (root+f'/locations/{region}/usages',{'api-version':'2024-10-01'})])
    prefetch_reads(api,initial)
    bundles=[];issues=[];total=len(configs)*len(configs[0]['subscription_ids'])
    indexes={}
    catalog_cache={};locations_cache={};defaults_done=0;defaults_total=0
    for config in configs:
        if not config.get('use_default'):continue
        for sid in config['subscription_ids']:
            if config['regions']=='auto':
                if sid not in locations_cache:locations_cache[sid]=[r['name'] for r in api.all(f'/subscriptions/{sid}/locations',{'api-version':'2022-12-01'}) if r.get('name') and re.fullmatch(r'[a-z0-9]{2,40}',r['name'])]
                defaults_total+=len(locations_cache[sid])
            else:defaults_total+=len(config['regions'])
    notify(0,defaults_total,'按区域确认 Azure 默认版本','默认版本')
    for config in configs:
        if not config.get('use_default'):continue
        regional={}
        for sid in config['subscription_ids']:
            if config['regions']=='auto':
                if sid not in locations_cache:locations_cache[sid]=[r['name'] for r in api.all(f'/subscriptions/{sid}/locations',{'api-version':'2022-12-01'}) if r.get('name') and re.fullmatch(r'[a-z0-9]{2,40}',r['name'])]
                candidates=locations_cache[sid]
            else:candidates=config['regions']
            regional[sid]={}
            for region in candidates:
                try:
                    key=(sid,region)
                    if key not in catalog_cache:catalog_cache[key]=api.all(f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices/locations/{region}/models')
                    regional[sid][region]=default_version(catalog_cache[key],config['model'],config['sku'])
                except (ValueError,RuntimeError):continue
                finally:
                    defaults_done+=1;notify(defaults_done,defaults_total,config['model']+' / '+region+' 默认版本已检查','默认版本')
            if not regional[sid]:issues.append(config['model']+'：所选订阅没有可确认的 Azure 默认版本，请指定版本后重新检查。')
        config['regional_versions']=regional
        config['model_version']=next((v for values in regional.values() for v in values.values()),'AzureDefault')
        config['subscription_versions']={sid:next(iter(values.values()),'AzureDefault') for sid,values in regional.items()}
    numbering_done=0;notify(0,total,'建立跨模型一致的区域编号','区域容量')
    for sid in configs[0]['subscription_ids']:
        order=list(REGIONS) if any(c['model']=='gpt-image-2' for c in configs) else []
        for config in configs:
            if config['regions']=='auto':
                try:
                    capacities=api.all(f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices/modelCapacities',{'api-version':batch.VERSION,'modelFormat':'OpenAI','modelName':config['model'],'modelVersion':config['model_version']})
                    discovered=sorted({row['location'].replace(' ','').lower() for row in capacities if row.get('properties',{}).get('skuName')==config['sku']})
                    discovered.extend(r for r in config.get('regional_versions',{}).get(sid,{}) if r not in discovered)
                except (ValueError,RuntimeError):discovered=[]
            else:discovered=config['regions']
            order.extend(region for region in discovered if region not in order)
            numbering_done+=1;notify(numbering_done,total,config['model']+' 区域容量已检查','区域容量')
        indexes[sid]={region:i for i,region in enumerate(order,1)}
    for config in configs:config['region_indexes']=indexes
    reads=[];parents={}
    for sid in configs[0]['subscription_ids']:
        root=f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices'
        try:
            for account in api.all(root+'/accounts'):
                if account.get('kind') in ('AIServices','OpenAI'):
                    parents[account['id'].lower()]=account
                    reads.append((account['id']+'/deployments',None))
                    if account.get('kind')=='AIServices' and account.get('properties',{}).get('allowProjectManagement'):
                        reads.append((account['id']+'/projects',None))
        except (ValueError,RuntimeError):pass
    prefetch_reads(api,reads)
    reads=[]
    for sid in configs[0]['subscription_ids']:
        root=f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices'
        for config in configs:
            try:
                occupied={account['location'].replace(' ','').lower() for account in parents.values()
                          if account['id'].split('/')[2].lower()==sid.lower()
                          and any(batch.same_model(d,config['model']) for d in api.all(account['id']+'/deployments'))}
                if isinstance(config['regions'],list) and set(config['regions'])<=occupied:continue
            except (ValueError,RuntimeError):continue  # Main scan reports incomplete inventory; no creation.
            versions=config.get('regional_versions',{}).get(sid,{})
            for version in dict.fromkeys(versions.values() if versions else [config['model_version']]):
                reads.append((root+'/modelCapacities',{'api-version':batch.VERSION,'modelFormat':'OpenAI','modelName':config['model'],'modelVersion':version}))
    prefetch_reads(api,reads)
    notify(0,total,'检查现有部署和配额','资源检查')
    for i,config in enumerate(configs):
        plan=batch.scan(api,config,on_progress=lambda done,count,label:notify(i*count+done,total,config['model']+' · '+label))
        authoritative=copy.deepcopy(plan);projects={};project_errors=[]
        for row in plan['rows']:
            if row['status'] not in ('existing','planned'):continue
            if row['status']=='existing':
                try:
                    quota=batch.model_metadata(api,row['subscription_id'],row['region'],{**config,'model_version':row['version']})
                    row.update(quota_checked=True,quota_remaining=quota['remaining'])
                except (ValueError,RuntimeError) as error:
                    row.update(quota_checked=False)
                    row['message']+=' 新建配额未确认：'+str(error)+'；已有部署仍只读复用。'
            aids=[d['account_id'] for d in row['existing']] if row['status']=='existing' else [row['account_id']]
            for aid in dict.fromkeys(aids):
                if row['status']=='existing':
                    projects[aid]={'name':'','new':False}
                    parent=parents.get(aid.lower()) or api.get(aid)
                    if parent.get('kind')=='AIServices' and parent.get('properties',{}).get('allowProjectManagement'):
                        try:
                            try:name=project_name(template,row,aid)
                            except ValueError:name=''
                            available=[p for p in api.all(aid+'/projects') if p.get('properties',{}).get('provisioningState')=='Succeeded']
                            matched=[p for p in available if p.get('name','').rsplit('/',1)[-1]==name]
                            chosen=matched or (available if len(available)==1 else [])
                            if chosen:projects[aid]['name']=chosen[0]['name'].rsplit('/',1)[-1]
                        except (ValueError,RuntimeError):
                            row['message']+=' Project 信息读取未完成；不创建或修改 Project。'
                    continue
                name=project_name(template,row,aid);path=aid+'/projects/'+name
                new_parent=row.get('new_account',False)
                if not new_parent:
                    parent=parents.get(aid.lower()) or api.get(aid)
                    if parent.get('kind')!='AIServices' or not parent.get('properties',{}).get('allowProjectManagement') or not parent.get('identity',{}).get('principalId'):
                        issues.append(config['model']+' · '+row['subscription']+' / '+row['region']+'：现有账户不支持 Project；不会自动改造。')
                        project_errors.append('现有资源不支持 Project；跳过此模型配置。')
                        row.update(status='blocked',message='现有资源未启用 Project 管理或托管身份；需人工处理。');break
                    present=api.optional(path)
                    if present is not None and present.get('properties',{}).get('provisioningState')!='Succeeded':
                        project_errors.append('现有 Project 尚未配置成功；跳过此模型配置。')
                        row.update(status='blocked',message='现有 Project 尚未配置成功。');break
                else:present=None
                projects[aid]={'name':name,'new':present is None}
        # Preserve the authoritative plan for full revalidation. Project blocks
        # apply only to the affected subscription/region, not the whole model.
        bundles.append({'plan':plan,'authoritative':authoritative,'projects':projects,'project_errors':project_errors})
        issues.extend(config['model']+'：'+error['message'] for error in plan['errors'])
        names_by_id={s['subscriptionId']:s.get('displayName',s['subscriptionId']) for s in session.subscriptions}
        for error in plan['errors']:
            sid=error['subscription_id']
            plan['rows'].append(dict(subscription_id=sid,subscription=names_by_id.get(sid,sid),region=error.get('region','—'),model=config['model'],version=config['model_version'],status='blocked',message=error['message']))
    if event.is_set():raise Cancelled()
    names={}
    for bundle in bundles:
        for row in bundle['plan']['rows']:
            if row.get('new_account') and row.get('account_id'):
                aid=row['account_id'];name=aid.rsplit('/',1)[-1].lower()
                if name in names and names[name].lower()!=aid.lower():raise ValueError('命名规则产生跨订阅的同名 Foundry；请改用包含 {sub8} 的规则，重新检查。')
                names[name]=aid
    return bundles,issues


def required_targets(row):
    deployments=row.get('existing',[]) if row['status']=='existing' else [{'account_id':row.get('account_id',''),'deployment':row.get('deployment','')}]
    return [(row['model'],row['region'],d['account_id'].lower(),d['deployment']) for d in deployments]


def start_scan(session,data):
    if session.api is None:raise ValueError('请先登录客户 Azure 账号。')
    configs,template=configurations(session,data)
    with session.lock:session.plan_id='';session.plans=[]
    def scan(event):
        selected=configs[0]['subscription_ids'];merged={};issues=[];done=0
        progress(session,0,len(selected),f'读取本次检查的订阅范围，最多 {MAX_CHECK_READS} 项只读请求并行','订阅检查')
        directory=job_api(session,event,snapshot=True)
        batch.subscriptions(directory,configs[0])
        def one(sid):
            if event.is_set():raise Cancelled()
            subset=copy.deepcopy(configs)
            for config in subset:config['subscription_ids']=[sid]
            def notify(*args,**kwargs):
                message=kwargs.get('message') or (args[2] if len(args)>2 else '')
                if message:progress(session,message=message)
            return scan_scope(session,subset,template,event,notify,directory.read_cache)
        with ThreadPoolExecutor(max_workers=MAX_CHECK_SUBSCRIPTIONS) as pool:
            futures={pool.submit(one,sid):sid for sid in selected}
            for future in as_completed(futures):
                if event.is_set():raise Cancelled()
                bundles,notes=future.result();issues.extend(notes)
                for bundle in bundles:
                    model=bundle['plan']['config']['model']
                    if model not in merged:
                        merged[model]=bundle
                    else:
                        target=merged[model]
                        for key in ('plan','authoritative'):
                            target[key]['rows'].extend(bundle[key]['rows']);target[key]['errors'].extend(bundle[key]['errors'])
                            for mapping in ('regional_versions','region_indexes','subscription_versions'):target[key]['config'].setdefault(mapping,{}).update(bundle[key]['config'].get(mapping,{}))
                        target['projects'].update(bundle['projects']);target['project_errors'].extend(bundle['project_errors'])
                done+=1;progress(session,done,len(selected),f'已检查 {done} / {len(selected)} 个订阅','订阅检查')
        bundles=[merged[c['model']] for c in configs];order={sid:i for i,sid in enumerate(selected)};names={}
        for bundle in bundles:
            for key in ('plan','authoritative'):
                bundle[key]['config']['subscription_ids']=list(selected)
                bundle[key]['rows'].sort(key=lambda r:order[r['subscription_id']])
                bundle[key]['errors'].sort(key=lambda r:order[r['subscription_id']])
            for row in bundle['plan']['rows']:
                if row.get('new_account') and row.get('account_id'):
                    aid=row['account_id'];name=aid.rsplit('/',1)[-1].lower()
                    if name in names and names[name].lower()!=aid.lower():raise ValueError('命名规则产生跨订阅的同名 Foundry；请使用含 {sub8} 的规则。')
                    names[name]=aid
        with session.lock:
            if event.is_set():raise Cancelled()
            session.plans=bundles;session.plan_id=secrets.token_urlsafe(24);session.issues=issues
            session.results=[];session.keys={}
            session.selected_subscriptions=[]
            for sub in session.subscriptions:
                if sub['subscriptionId'] not in selected:continue
                targets={target for b in bundles for r in b['plan']['rows'] if r['subscription_id']==sub['subscriptionId'] for target in required_targets(r)}
                session.selected_subscriptions.append({**sub,'expected_models':len(targets),'expected_targets':[list(t) for t in sorted(targets)]})
            session.job['partial']=any(r['status']=='blocked' for b in bundles for r in b['plan']['rows']) or bool(issues)
    start(session,'检查现有部署、Project 与实际配额',scan)


def start_catalog(session,data):
    if session.api is None:raise ValueError('请先登录客户 Azure 账号。')
    sid=str(UUID(data.get('subscription_id','')));region=data.get('region','')
    if sid not in {s['subscriptionId'] for s in session.subscriptions} or not re.fullmatch(r'[a-z0-9]{2,40}',region):raise ValueError('请选择当前账号的订阅和合法区域。')
    def catalog(event):
        api=JobAzure(session.credential,event,lambda message:progress(session,message=message))
        rows=api.all(f'/subscriptions/{sid}/providers/Microsoft.CognitiveServices/locations/{region}/models')
        available={name:[] for name in batch.MODELS}
        for row in rows:
            model=row.get('model') or row
            if model.get('format')=='OpenAI' and model.get('name') in available:
                available[model['name']].append({'version':model.get('version'),'default':model.get('isDefaultVersion') is True,'skus':[s.get('name') for s in model.get('skus',[])]})
        if event.is_set():raise Cancelled()
        with session.lock:
            if event.is_set() or session.api is None:raise Cancelled()
            session.catalog={'subscription_id':sid,'region':region,'models':available}
    start(session,'读取模型目录',catalog)


def reuse_existing(api,expected,on_progress):
    """Read exact server-reviewed deployments; this path cannot create resources."""
    report={'tenant_id':expected['tenant_id'],'rows':[],'errors':[]}
    for row in expected['rows']:
        item={**row,'outputs':[]};report['rows'].append(item)
        try:
            for deployed in row['existing']:
                aid=deployed['account_id']
                if aid.split('/')[2].lower()!=row['subscription_id'].lower():raise ValueError('复用资源不属于清单订阅。')
                account=api.get(aid);actual=api.get(aid+'/deployments/'+deployed['deployment'])
                if account.get('properties',{}).get('provisioningState')!='Succeeded' or actual.get('properties',{}).get('provisioningState')!='Succeeded':raise ValueError('已有资源或部署尚未成功，请重新检查。')
                output=batch.account_result(account,actual)
                if not batch.same_model(actual,row['model'],deployed['sku']) or output['capacity']!=deployed['capacity'] or output['model'].get('version')!=deployed['version']:raise ValueError('已有部署状态已变化，请重新检查。')
                item['outputs'].append(output)
        except (batch.AzureError,ValueError,RuntimeError) as error:item.update(status='failed',message=str(error))
        finally:on_progress(report)
    return report


def start_deploy(session,data):
    if session.api is None:raise ValueError('请先登录。')
    if not session.plan_id or data.get('plan_id')!=session.plan_id or data.get('confirmed') is not True:raise ValueError('请先检查并确认当前服务端清单；不能直接上传创建清单。')
    bundles=copy.deepcopy(session.plans)
    if not any(any(r['status'] in ('planned','existing') for r in b['plan']['rows']) for b in bundles):raise ValueError('清单没有可执行项，请调整后重新检查。')
    def deploy(event):
        if not WRITE_LOCK.acquire(blocking=False):raise ValueError('另一项创建任务正在执行，请稍后重新检查。')
        try:deploy_locked(event)
        finally:WRITE_LOCK.release()
    def deploy_locked(event):
        selected=list(dict.fromkeys(r['subscription_id'] for b in bundles for r in b['plan']['rows']))
        result={};issues=list(session.issues);done=0
        progress(session,0,len(selected),'最多同时处理三个订阅；每个订阅内依次执行','部署与 Project')
        with session.lock:session.plan_id='';session.results=[];session.keys={}
        def publish(row,output,project,status):
            value=result_row(row,output,project,status)
            with session.lock:
                result[(output['account_id'],output['deployment'])]=value
                if session.api is not None:session.results=copy.deepcopy(list(result.values()))
        def note(message):
            with session.lock:issues.append(message);session.issues=list(issues)
        def one(sid):
            api=job_api(session,event);owned_accounts=set()
            for bundle in bundles:
                if event.is_set():raise Cancelled()
                expected=copy.deepcopy(bundle.get('authoritative',bundle['plan']))
                expected['config']['subscription_ids']=[sid]
                if sid in expected['config'].get('subscription_versions',{}):expected['config']['model_version']=expected['config']['subscription_versions'][sid]
                expected['rows']=[r for r in expected['rows'] if r['subscription_id']==sid]
                expected['errors']=[e for e in expected['errors'] if e['subscription_id']==sid]
                if not expected['rows']:continue
                blocked={(r['subscription_id'],r['region']):r.get('message') or '检查未通过。' for r in bundle['plan']['rows'] if r['subscription_id']==sid and r['status']=='blocked'}
                planned={(r['subscription_id'],r['region'],r['model']) for r in expected['rows'] if r['status']=='planned'}
                for row in expected['rows']:
                    if row.get('new_account') and row['account_id'] in owned_accounts:row['new_account']=False
                def checkpoint(current):
                    for row in current['rows']:
                        for output in row.get('outputs',[]):publish(row,output,'','部署已就绪；等待最终核验')
                try:
                    if not blocked and not expected['errors'] and all(r['status']=='existing' and r.get('existing') for r in expected['rows']):report=reuse_existing(api,expected,checkpoint)
                    else:report=batch.execute(api,expected,on_progress=checkpoint,**({'blocked':blocked} if blocked else {}))
                except (batch.AzureError,ValueError,RuntimeError) as error:
                    note(sid+' / '+expected['config'].get('model','')+'：'+str(error));continue
                for row in report['rows']:
                    if row.get('status') in ('failed','blocked'):note(row['subscription']+' / '+row['region']+'：'+row.get('message','配置失败'))
                    for output in row.get('outputs',[]):
                        aid=output['account_id'];project=bundle['projects'][aid]['name'];project_state=''
                        owned_accounts.add(aid)
                        try:
                            if (row['subscription_id'],row['region'],row['model']) in planned:ensure_project(api,aid,project,row['region'])
                        except (batch.AzureError,ValueError,RuntimeError):
                            project_state='Project 配置未完成；请重新检查。';note(row['subscription']+' / '+row['region']+'：'+project_state)
                        publish(row,output,project if not project_state else '',project_state or row['status'])
        with ThreadPoolExecutor(max_workers=MAX_SUBSCRIPTIONS_IN_FLIGHT) as pool:
            def guarded(sid):
                try:return one(sid)
                except Exception:
                    event.set();raise
            futures=[pool.submit(guarded,sid) for sid in selected]
            for future in as_completed(futures):
                future.result();done+=1;progress(session,done,len(selected),f'已处理 {done} / {len(selected)} 个订阅','部署与 Project')
        verify_results(session,event)
        collect_keys(session,event)
        with session.lock:
            session.issues=list(dict.fromkeys(session.issues))
            session.job['partial']=session.job.get('partial',False) or bool(session.issues)
    start(session,'创建缺失项、核验实际部署并准备 Excel',deploy)


def verify_results(session,event):
    groups={}
    for row in session.results:groups.setdefault(row['_account_id'],[]).append(row)
    total=len(session.results);done=0
    progress(session,0,total,'重新读取 Azure 的资源、Project 和模型状态','最终核验')
    def one(aid,rows):
        api=job_api(session,event);projects={}
        try:account=api.get(aid)
        except (batch.AzureError,ValueError,RuntimeError) as error:
            with session.lock:
                for row in rows:row.update(verified=False,status='资源核验未完成：'+str(error))
                session.issues.append(aid.rsplit('/',1)[-1]+'：'+str(error));session.job['partial']=True
            return
        for row in rows:
            try:
                actual=api.get(aid+'/deployments/'+row['deployment'])
                props=actual.get('properties',{});model=props.get('model',{});sku=actual.get('sku',{})
                if account.get('properties',{}).get('provisioningState')!='Succeeded' or props.get('provisioningState')!='Succeeded':raise ValueError('Azure 资源或模型尚未 Succeeded。')
                if model.get('format')!='OpenAI' or not model.get('version') or not sku.get('name') or not isinstance(sku.get('capacity'),(int,float)) or sku['capacity']<=0:raise ValueError('Azure 实际部署元数据不完整，不能确认版本或容量。')
                if model.get('name')!=row['model'] or model.get('version')!=row['version'] or sku.get('name')!=row['sku'] or sku.get('capacity')!=row['capacity']:raise ValueError('实际模型、版本、类型或容量与结果不一致。')
                if row.get('region_code') and account.get('location','').replace(' ','').lower()!=row['region_code']:raise ValueError('实际 Foundry 区域与结果不一致。')
                if row['资源']:
                    if row['资源'] not in projects:projects[row['资源']]=api.get(aid+'/projects/'+row['资源'])
                    if projects[row['资源']].get('properties',{}).get('provisioningState')!='Succeeded':raise ValueError('Project 尚未 Succeeded。')
                elif account.get('kind')=='AIServices':raise ValueError('Project 未确认；模型已存在但交付信息不完整。')
                with session.lock:row.update(verified=True,status='已核验',rpm=batch.rpm(props.get('rateLimits')),tpm=batch.tpm(props.get('rateLimits')))
            except (batch.AzureError,ValueError,RuntimeError) as error:
                with session.lock:
                    row.update(verified=False,status='核验未通过：'+str(error));session.issues.append(row['订阅']+' / '+row['model']+'：'+str(error));session.job['partial']=True
    with ThreadPoolExecutor(max_workers=MAX_SUBSCRIPTIONS_IN_FLIGHT) as pool:
        futures={pool.submit(one,aid,rows):rows for aid,rows in groups.items()}
        for future in as_completed(futures):
            future.result();done+=len(futures[future]);progress(session,done,total,f'已核验 {done} / {total} 个实际模型部署','最终核验')


def collect_keys(session,event):
    aids=sorted({r['_account_id'] for r in session.results if r.get('verified')})
    progress(session,0,len(aids),'只读取当前 Key，不重置密钥','准备模板导出')
    def one(aid):
        api=job_api(session,event)
        try:
            account=api.get(aid)
            if account.get('properties',{}).get('disableLocalAuth'):raise ValueError('已禁用 Key 认证，保持安全策略。')
            key=api.call('POST',aid+'/listKeys',params={'api-version':batch.VERSION}).get('key1')
            if not key:raise ValueError('Azure 未返回 Key，请核对 listKeys 权限。')
            with session.lock:
                if event.is_set() or session.api is None:raise Cancelled()
                session.keys[aid]=key
        except (batch.AzureError,ValueError,RuntimeError) as error:
            with session.lock:session.issues.append(aid.rsplit('/',1)[-1]+'：'+str(error));session.job['partial']=True
    with ThreadPoolExecutor(max_workers=MAX_SUBSCRIPTIONS_IN_FLIGHT) as pool:
        futures=[pool.submit(one,aid) for aid in aids]
        for i,future in enumerate(as_completed(futures),1):
            future.result();progress(session,i,len(aids),f'已处理 {i} / {len(aids)} 个资源 Key','准备模板导出')


def result_row(row,output,project,status):
    aid=output['account_id'];actual=output.get('model') or {}
    endpoint=output.get('endpoints',{}).get('AI Foundry API') or output.get('endpoint') or ''
    route='/openai/v1/images/generations' if (actual.get('name') or row['model']).startswith('gpt-image-') else '/openai/v1/responses'
    return {'订阅':row['subscription'],'订阅 ID':row['subscription_id'],'资源组':aid.split('/')[4],'资源名称-foundry':aid.rsplit('/',1)[-1],'区域':REGIONS.get(row['region'],row['region']),'资源':project,'url':endpoint.rstrip('/')+route if endpoint else '', '密钥':'未读取','_account_id':aid,'model':actual.get('name') or row['model'],'version':actual.get('version') or '', 'deployment':output['deployment'],'sku':output.get('sku',row.get('sku','')), 'capacity':output['capacity'],'rpm':output.get('rpm'),'tpm':output.get('tpm'),'status':status,'verified':False,'region_code':row['region']}


def read_keys(session):
    if session.api is None or not session.results:raise ValueError('请先完成检查和创建，获得实际资源结果。')
    def read(event):
        api=JobAzure(session.credential,event,lambda message:progress(session,message=message))
        aids=sorted({row['_account_id'] for row in session.results})
        for i,aid in enumerate(aids):
            account=api.get(aid)
            if account.get('properties',{}).get('disableLocalAuth'):
                with session.lock:session.issues.append(aid.rsplit('/',1)[-1]+'：禁用 Key 认证，未读取密钥。')
            else:
                keys=api.call('POST',aid+'/listKeys',params={'api-version':batch.VERSION})
                key=keys.get('key1')
                if not key:raise ValueError('Azure 未返回账户 Key；请检查 listKeys 权限。')
                with session.lock:
                    if event.is_set() or session.api is None:raise Cancelled()
                    session.keys[aid]=key
            progress(session,i+1,len(aids),'已读取 '+str(i+1)+' / '+str(len(aids))+' 个资源')
    start(session,'读取资源 Key，不重置密钥',read)


def excel(rows,details=None,subscriptions=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Sheet1'
    def natural(value):return tuple(int(s) if s.isdigit() else s.lower() for s in re.split(r'(\d+)',str(value)))
    def order(row):return (natural(row.get('订阅','')),natural(row.get('资源名称-foundry','')),row.get('model',''),row.get('deployment',''))
    def fill(sheet,columns,data):
        sheet.append(list(columns))
        for row in data:
            sheet.append([row.get(k,'') if row.get(k) is not None else '' for k in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        for cell in sheet[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='0078D4')
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for i,column in enumerate(columns,1):sheet.column_dimensions[get_column_letter(i)].width=44 if column in ('url','密钥','订阅 ID') else 28
    # Rebuild a clean template; never import customer hyperlinks, keys or metadata.
    unique={}
    for row in rows:
        key=row.get('_account_id') or tuple(str(row.get(k,'')) for k in COLUMNS[:6])
        candidate={**row,'资源名称-foundry项目':row.get('资源名称-foundry','')}
        if key not in unique or '/images/generations' in row.get('url',''):unique[key]=candidate
    fill(sheet,TEMPLATE_COLUMNS,sorted(unique.values(),key=order))
    fill(book.create_sheet('Sheet2'),('订阅','订阅 ID','资源名称-foundry','区域','资源','model','version','sku','deployment','capacity','rpm','tpm','status','url'),sorted(details or rows,key=order))
    scopes=subscriptions or list({r['订阅 ID']:{'subscriptionId':r['订阅 ID'],'displayName':r['订阅']} for r in rows}.values())
    summary=[]
    for sub in sorted(scopes,key=lambda s:natural(s.get('displayName',s['subscriptionId']))):
        found=[r for r in rows if r['订阅 ID']==sub['subscriptionId']]
        expected=sub.get('expected_models',len(found));verified=sum(r.get('verified') is True for r in found)
        if 'expected_targets' in sub:
            targets={tuple(t) for t in sub['expected_targets']};expected=len(targets)
            verified=len(targets & {(r.get('model'),r.get('region_code'),r.get('_account_id','').lower(),r.get('deployment')) for r in found if r.get('verified') is True})
        summary.append({'订阅':sub.get('displayName',sub['subscriptionId']),'订阅 ID':sub['subscriptionId'],'Foundry 数量':len({r.get('_account_id') or r.get('资源名称-foundry') for r in found}),'Project 数量':len({(r.get('_account_id'),r['资源']) for r in found if r.get('资源')}),'计划模型数量':expected,'返回模型数量':len(found),'已核验模型':verified,'未完成项':max(0,expected-verified),'已读取 Key':sum(bool(r.get('密钥') and r['密钥']!='未读取') for r in unique.values() if r['订阅 ID']==sub['subscriptionId']),'结果':'已核验' if expected and verified>=expected and all(r.get('verified') is True for r in found) else '未完成核验'})
    fill(book.create_sheet('Sheet3'),tuple(summary[0]) if summary else ('订阅','订阅 ID','结果'),summary)
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def export_details(session):
    rows=copy.deepcopy(session.results)
    returned={(r['订阅 ID'],r['model'],r.get('region_code'),r.get('_account_id','').lower(),r.get('deployment')) for r in rows}
    for bundle in session.plans:
        for row in bundle['plan']['rows']:
            for target in required_targets(row):
                if (row['subscription_id'],*target) in returned:continue
                aid=target[2]
                rows.append({'订阅':row['subscription'],'订阅 ID':row['subscription_id'],'资源名称-foundry':aid.rsplit('/',1)[-1],'区域':REGIONS.get(row['region'],row['region']),'model':row['model'],'version':row.get('version',''),'sku':row.get('sku',''),'deployment':target[3],'capacity':row.get('capacity'),'status':'未返回完整结果；'+row.get('message','请重新检查。')})
    return rows


class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args):pass

    def reply(self,status,data,kind='application/json; charset=utf-8',cookie=None,download=False):
        if isinstance(data,dict):data=json.dumps(data,ensure_ascii=False).encode()
        self.send_response(status);self.send_header('Content-Type',kind);self.send_header('Content-Length',str(len(data)))
        self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('Referrer-Policy','no-referrer')
        self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
        if cookie:self.send_header('Set-Cookie',cookie)
        if download:self.send_header('Content-Disposition','attachment; filename="foundry-resources.xlsx"')
        self.end_headers();self.wfile.write(data)

    def session(self):
        cookie=SimpleCookie();cookie.load(self.headers.get('Cookie',''))
        sid=cookie.get('foundry_session');sid=sid.value if sid else ''
        with self.server.sessions_lock:
            for value in self.server.sessions.values():
                with value.lock:
                    if value.oauth and value.oauth['expires']<time.monotonic():value.oauth=None;value.login_url=''
            expired=[key for key,value in self.server.sessions.items() if time.monotonic()-value.touched>14400]
            for key in expired:self.server.sessions.pop(key).cancel.set()
            session=self.server.sessions.get(sid)
            if session:session.touched=time.monotonic()
            return session

    def trusted_host(self):
        if PUBLIC_ORIGIN:return self.headers.get('Host','')==urlsplit(PUBLIC_ORIGIN).netloc
        try:return urlsplit('http://'+self.headers.get('Host','')).hostname in ('127.0.0.1','localhost','::1')
        except ValueError:return False

    def do_GET(self):
        if urlsplit(self.path).path=='/healthz':return self.reply(200,{'status':'ok','version':WEB_VERSION})
        if not self.trusted_host():return self.reply(403,{'error':'请通过 localhost 或 SSH 隧道访问此客户本地容器。'})
        path=urlsplit(self.path).path
        if path=='/healthz':return self.reply(200,{'status':'ok'})
        if path=='/':
            session=self.session();cookie=None
            if session is None:
                sid=secrets.token_urlsafe(32);session=Session()
                with self.server.sessions_lock:
                    if len(self.server.sessions)>=256:
                        unused=sorted((key for key,value in self.server.sessions.items() if value.api is None and value.job.get('status') not in ACTIVE),key=lambda key:self.server.sessions[key].touched)
                        for key in unused[:128]:self.server.sessions.pop(key).cancel.set()
                    if len(self.server.sessions)>=256:return self.reply(503,{'error':'当前活动会话已满，请稍后重试。'})
                    self.server.sessions[sid]=session
                cookie='foundry_session='+sid+'; Path=/; HttpOnly; SameSite=Lax'+('; Secure' if PUBLIC_ORIGIN else '')
            return self.reply(200,(ROOT/'web/index.html').read_bytes(),'text/html; charset=utf-8',cookie)
        files={'/app.js':'text/javascript; charset=utf-8','/style.css':'text/css; charset=utf-8'}
        if path in files:return self.reply(200,(ROOT/'web'/path[1:]).read_bytes(),files[path])
        session=self.session()
        if path=='/auth/callback' and session:
            try:
                query=parse_qs(urlsplit(self.path).query)
                if any(len(values)!=1 for values in query.values()):raise ValueError('授权响应格式无效。')
                finish_browser_login(session,{key:values[0] for key,values in query.items()})
            except ValueError as error:
                with session.lock:session.login_error=str(error)
            except Exception:
                with session.lock:session.login_error='授权回调未完成，请重新登录。'
            self.send_response(303);self.send_header('Location','/');self.send_header('Cache-Control','no-store');self.send_header('Referrer-Policy','no-referrer');self.send_header('Content-Length','0');self.end_headers();return
        if path=='/api/state' and session:return self.reply(200,session.snapshot())
        return self.reply(404,{'error':'页面或接口不存在。'})

    def do_POST(self):
        expected=PUBLIC_ORIGIN or 'http://'+self.headers.get('Host','')
        if not self.trusted_host() or self.headers.get('Origin')!=expected or self.headers.get('X-Requested-With')!='FoundryUI':return self.reply(403,{'error':'仅接受工具网页发起的同源请求。'})
        session=self.session()
        if session is None:return self.reply(401,{'error':'会话已过期，请重新打开网页。'})
        try:
            length=int(self.headers.get('Content-Length','0'))
            if not 0<length<=65536 or self.headers.get('Content-Type')!='application/json':raise ValueError('请求大小或内容格式无效。')
            data=json.loads(self.rfile.read(length))
            if not isinstance(data,dict):raise ValueError('请求格式无效。')
            path=urlsplit(self.path).path
            if path=='/api/login':start_login(session,data)
            elif path=='/api/catalog':start_catalog(session,data)
            elif path=='/api/scan':start_scan(session,data)
            elif path=='/api/deploy':start_deploy(session,data)
            elif path=='/api/keys':read_keys(session)
            elif path=='/api/cancel':
                cancel_job(session)
            elif path=='/api/logout':
                session.cancel.set()
                with session.lock:session.api=None;session.credential=None;session.subscriptions=[];session.plans=[];session.plan_id='';session.results=[];session.keys={};session.challenge=None;session.catalog={};session.issues=[];session.oauth=None;session.login_url='';session.login_generation+=1
            elif path=='/api/reveal':
                key=session.keys.get(data.get('account_id'))
                if not key:raise ValueError('此资源尚未读取密钥；先点击读取 Key。')
                return self.reply(200,{'key':key})
            elif path=='/api/export':
                if not session.results:raise ValueError('尚无实际资源结果可导出。')
                include=data.get('include_keys') is True
                if include and data.get('confirmed') is not True:raise ValueError('请确认导出文件包含明文密钥。')
                rows=[{**r,'密钥':session.keys.get(r['_account_id'],'未读取') if include else ''} for r in session.results]
                return self.reply(200,excel(rows,export_details(session),session.selected_subscriptions),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',download=True)
            else:return self.reply(404,{'error':'接口不存在。'})
            cookie=None
            if path=='/api/login' and os.environ.get('OAUTH_CLIENT_ID'):
                values=SimpleCookie();values.load(self.headers.get('Cookie',''))
                cookie='foundry_session='+values['foundry_session'].value+'; Path=/; HttpOnly; SameSite=Lax'+('; Secure' if PUBLIC_ORIGIN else '')
            return self.reply(202,{'ok':True},cookie=cookie)
        except (ValueError,KeyError,TypeError):return self.reply(400,{'error':'输入或状态无效，请检查选择、版本及命名规则后重试。'})
        except Exception:return self.reply(500,{'error':'请求未完成，请重新检查；不要反复提交创建。'})


def make_server(host='127.0.0.1',port=8767):
    if host not in ('127.0.0.1','0.0.0.0'):raise ValueError('仅支持本地监听或容器内部监听。')
    server=ThreadingHTTPServer((host,port),Handler);server.daemon_threads=True
    server.sessions={};server.sessions_lock=threading.RLock();return server


def main(argv=None):
    parser=argparse.ArgumentParser(description='客户浏览器工作台，不依赖 Azure CLI。')
    parser.add_argument('--host',default='127.0.0.1');parser.add_argument('--port',type=int,default=8767)
    args=parser.parse_args(argv)
    # SDK authentication failures must not emit raw token/tenant diagnostics.
    logging.getLogger('azure').setLevel(logging.CRITICAL)
    logging.getLogger('msal').setLevel(logging.CRITICAL)
    server=make_server(args.host,args.port)
    print(f'Foundry browser workspace: http://127.0.0.1:{args.port} (SSH tunnel for remote Ubuntu)',flush=True)
    try:server.serve_forever()
    except KeyboardInterrupt:pass
    finally:server.server_close()


if __name__=='__main__':main()
